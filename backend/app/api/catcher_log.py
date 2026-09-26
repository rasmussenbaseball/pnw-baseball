"""Coach catcher quick log (xlsx) for the TrackMan Suite's Catching tab.

A coach charts every live game by hand in a small workbook ("Bushnell
Catcher Quick Log"). The CURRENT format (Sept 2026) has two sheets:

  Live Log    Inning, Catcher, Batter, Pitcher, Count, Event, Outcome,
              Two Strike?, Foul Balls, Optional Note(s)
  Throwdowns  Inning, Catcher, Pop Time, Result, Notes   (practice throws)

Columns are read BY HEADER NAME, so the older layout (Count in column C,
Pitcher Time, GM Grade) still parses, and future columns can be added
without breaking the upload. TrackMan sees none of this (no runners, no
blocks), so the log is the other half of the catcher picture; block events
are matched back to the exact TrackMan pitch by catcher + pitcher + inning
+ count when one pitch fits.

Parsing is forgiving: Excel turns counts like 2-2 into dates (Feb 2),
catchers are typed "First Last" or "Nick / First Last", event and outcome
spellings drift ("Clank", "Drop third whiff", "Block and throw"). Games in
one sheet split where the inning number resets. Re-uploading a date
replaces that date's rows.
"""
import io
import re
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from psycopg2.extras import execute_values

from ..models.database import get_connection
from .trackman_suite import _gate, _write_gate

router = APIRouter()

BLOCK_OUTCOMES = {"Block Success", "Block Miss", "Passed Ball"}
THROW_OUTCOMES = {"Stolen Base", "Caught Stealing", "Pickoff"}
OUTCOMES = [
    (re.compile(r"block\s*succ|blocked|success", re.I), "Block Success"),
    (re.compile(r"block\s*miss|\bmiss", re.I), "Block Miss"),
    (re.compile(r"passed|\bpb\b", re.I), "Passed Ball"),
    (re.compile(r"caught\s*steal|\bcs\b", re.I), "Caught Stealing"),
    (re.compile(r"stolen|\bsb\b", re.I), "Stolen Base"),
    (re.compile(r"pick", re.I), "Pickoff"),
    (re.compile(r"throw\s*down|throwdown", re.I), "Throwdown"),
    (re.compile(r"\bout\b", re.I), "Out"),
    (re.compile(r"safe", re.I), "Safe"),
    (re.compile(r"caught", re.I), "Caught"),
]
POP_RE = re.compile(r"\b([12]\.\d{1,2})\b")
_COLS = ("game_no", "row_no", "inning", "catcher", "batter", "pitcher", "count", "event", "outcome",
         "two_strike", "foul_balls", "pop_time", "grade", "note")


def _ensure_table(cur):
    cur.execute("""
        CREATE TABLE IF NOT EXISTS tm_catcher_log (
            id            SERIAL PRIMARY KEY,
            owner_user_id UUID NOT NULL,
            session_date  DATE NOT NULL,
            game_no       INTEGER NOT NULL DEFAULT 1,
            row_no        INTEGER,
            inning        INTEGER,
            catcher       TEXT NOT NULL,
            count         TEXT,
            event         TEXT,
            outcome       TEXT,
            pop_time      DOUBLE PRECISION,
            grade         DOUBLE PRECISION,
            note          TEXT,
            created_at    TIMESTAMPTZ DEFAULT now()
        )
    """)
    for col, typ in (("batter", "TEXT"), ("pitcher", "TEXT"), ("two_strike", "BOOLEAN"), ("foul_balls", "INTEGER")):
        cur.execute(f"ALTER TABLE tm_catcher_log ADD COLUMN IF NOT EXISTS {col} {typ}")
    cur.execute("ALTER TABLE tm_catcher_log ENABLE ROW LEVEL SECURITY")
    cur.execute("CREATE INDEX IF NOT EXISTS tm_catcher_log_owner_date ON tm_catcher_log (owner_user_id, session_date)")


def _count(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return f"{v.month}-{v.day}"
    m = re.match(r"^(\d)\s*[-–/]\s*(\d)$", str(v).strip())
    return f"{m.group(1)}-{m.group(2)}" if m else None


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _name(v):
    """'Hoko / Reyn Gaspar' -> 'Reyn Gaspar'; trims stray spaces."""
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    if "/" in s:
        s = s.split("/")[-1].strip()
    return s or None


def _yes(v):
    return None if v is None else str(v).strip().lower() in ("yes", "y", "true", "1")


def _outcome(v):
    if not v:
        return None
    s = str(v).strip()
    for rx, label in OUTCOMES:
        if rx.search(s):
            return label
    return s


def _event(raw, outcome):
    """Normalize the free-text event into Block / Throw / Drop 3rd / Pop /
    Throwdown, letting the outcome settle ambiguous text ('Clank' with a
    passed ball is a block chance; 'Steal, wild pitch' is a throw)."""
    e = (raw or "").strip().lower()
    if "third" in e or "3rd" in e or "drop" in e:
        return "Drop 3rd"
    if "pop" in e:
        return "Pop"
    if "throwdown" in e or "throw down" in e:
        return "Throwdown"
    if outcome in BLOCK_OUTCOMES or e.startswith("block"):
        return "Block"
    if outcome in THROW_OUTCOMES or "throw" in e or "steal" in e:
        return "Throw"
    return e.title() if e else None


def _header_map(row):
    """{normalized header: column index} for a header row, or None."""
    cells = [str(c).strip().lower() if c is not None else "" for c in row]
    if "inning" not in cells or "catcher" not in cells:
        return None
    m = {}
    for i, c in enumerate(cells):
        if not c:
            continue
        key = re.sub(r"[^a-z]", "", c)
        m.setdefault(key, i)          # first "Optional Note" wins; the rest are gathered below
    m["_notes"] = [i for i, c in enumerate(cells) if "note" in c]
    return m


def _parse_live(ws):
    rows, hm, prev_inning, game_no, row_no = [], None, None, 1, 0
    for raw in ws.iter_rows(values_only=True):
        cells = list(raw)
        if hm is None:
            hm = _header_map(cells)
            continue
        get = lambda k: cells[hm[k]] if k in hm and hm[k] < len(cells) else None  # noqa: E731
        inning, catcher = _num(get("inning")), _name(get("catcher"))
        if not catcher:
            continue
        outcome = _outcome(get("outcome"))
        event = _event(get("event"), outcome)
        if not event and outcome is None:
            continue
        inning = int(inning) if inning is not None else None
        if prev_inning is not None and inning is not None and inning < prev_inning:
            game_no += 1
        if inning is not None:
            prev_inning = inning
        notes = [str(cells[i]).strip() for i in hm.get("_notes", []) if i < len(cells) and cells[i] is not None and str(cells[i]).strip()]
        note = " · ".join(dict.fromkeys(notes)) or None
        pop = _num(get("pitchertime")) or _num(get("poptime"))
        if pop is None and note:
            m = POP_RE.search(note)
            if m:
                pop = float(m.group(1))
        if pop is not None and not (1.5 <= pop <= 3.2):
            pop = None
        fouls = _num(get("foulballs"))
        row_no += 1
        rows.append({
            "game_no": game_no, "row_no": row_no, "inning": inning, "catcher": catcher,
            "batter": _name(get("batter")), "pitcher": _name(get("pitcher")),
            "count": _count(get("count")), "event": event, "outcome": outcome,
            "two_strike": _yes(get("twostrike")), "foul_balls": int(fouls) if fouls is not None else None,
            "pop_time": pop, "grade": _num(get("gmgrade")), "note": note,
        })
    return rows


def _parse_throwdowns(ws, start_row_no):
    rows, hm, row_no = [], None, start_row_no
    for raw in ws.iter_rows(values_only=True):
        cells = list(raw)
        if hm is None:
            hm = _header_map(cells)
            continue
        get = lambda k: cells[hm[k]] if k in hm and hm[k] < len(cells) else None  # noqa: E731
        catcher = _name(get("catcher"))
        if not catcher:
            continue
        pop = _num(get("poptime"))
        note = str(get("notes") or get("note") or "").strip() or None
        if pop is None and note:
            m = POP_RE.search(note)
            if m:
                pop = float(m.group(1))
        result = get("result")
        inning = _num(get("inning"))
        row_no += 1
        rows.append({
            "game_no": 0, "row_no": row_no, "inning": int(inning) if inning is not None else None,
            "catcher": catcher, "batter": None, "pitcher": None, "count": None,
            "event": "Throwdown", "outcome": _outcome(result) or (str(result).strip() if result else "Throwdown"),
            "two_strike": None, "foul_balls": None,
            "pop_time": pop if (pop is None or 1.5 <= pop <= 3.2) else None, "grade": None, "note": note,
        })
    return rows


def parse_workbook(data: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    live = next((w for w in wb.worksheets if "log" in w.title.lower()), wb.worksheets[0])
    rows = _parse_live(live)
    td = next((w for w in wb.worksheets if "throwdown" in w.title.lower()), None)
    if td is not None:
        rows += _parse_throwdowns(td, len(rows))
    if not rows:
        raise HTTPException(status_code=400, detail="No log rows found (expected a 'Live Log' sheet with an Inning / Catcher / Count / Event / Outcome header).")
    return rows


@router.post("/trackman/catcher-log/upload")
async def catcher_log_upload(
    file: UploadFile = File(...),
    session_date: str | None = Form(None),
    owner: str = Depends(_write_gate),
):
    sdate = session_date or date.today().isoformat()
    rows = parse_workbook(await file.read())
    with get_connection() as conn:
        cur = conn.cursor()
        _ensure_table(cur)
        cur.execute("DELETE FROM tm_catcher_log WHERE owner_user_id = %s AND session_date = %s", (owner, sdate))
        execute_values(cur, f"""
            INSERT INTO tm_catcher_log (owner_user_id, session_date, {", ".join(_COLS)})
            VALUES %s""",
            [tuple([owner, sdate] + [r[c] for c in _COLS]) for r in rows])
        conn.commit()
    games = max((r["game_no"] for r in rows), default=1)
    return {"status": "ok", "session_date": sdate, "rows": len(rows), "games": games,
            "throwdowns": sum(1 for r in rows if r["event"] == "Throwdown"),
            "catchers": sorted({r["catcher"] for r in rows})}


def _last(name):
    n = (name or "").strip()
    if not n:
        return ""
    return (n.split(",")[0] if "," in n else n.split()[-1]).lower()


def _new_agg(name):
    return {"catcher": name, "games": set(), "block_opps": 0, "block_success": 0, "block_miss": 0, "pb": 0,
            "k2_opps": 0, "k2_success": 0, "sb": 0, "cs": 0, "pickoffs": 0, "drop3": 0, "drop3_outs": 0,
            "throwdowns": 0, "pops": [], "grades": [], "notes": []}


@router.get("/trackman/catcher-log")
def catcher_log(date_from: str | None = Query(None), date_to: str | None = Query(None),
                owner: str = Depends(_gate)):
    """Per-catcher rollup of the coach log, blocking by pitcher, and every
    event with its TrackMan pitch matched (catcher + pitcher + inning +
    count, ball / dirt ball preferred) when exactly one fits."""
    with get_connection() as conn:
        cur = conn.cursor()
        _ensure_table(cur)
        extra, params = "", [owner]
        if date_from:
            extra += " AND session_date >= %s"; params.append(date_from)
        if date_to:
            extra += " AND session_date <= %s"; params.append(date_to)
        cur.execute(f"""SELECT * FROM tm_catcher_log WHERE owner_user_id = %s{extra}
                        ORDER BY session_date, game_no, row_no""", params)
        rows = [dict(r) for r in cur.fetchall()]
        if not rows:
            return {"catchers": [], "by_pitcher": [], "events": [], "dates": []}
        dates = sorted({r["session_date"] for r in rows})
        # A workbook holds several games and is uploaded under the last one's
        # date, so candidates come from the live sessions of the prior six
        # weeks; the match must still be unique.
        cur.execute("""SELECT p.id, s.session_date, p.catcher, p.pitcher, p.inning, p.balls, p.strikes, p.pitch_call,
                              COALESCE(p.override_pitch_type, p.class_pitch_type, p.tagged_pitch_type, p.auto_pitch_type) AS ptype,
                              p.rel_speed, p.plate_loc_side, p.plate_loc_height
                       FROM tm_pitches p JOIN tm_sessions s ON s.id = p.session_id
                       WHERE p.owner_user_id = %s AND p.catcher IS NOT NULL
                         AND s.session_type IN ('game', 'scrimmage', 'intrasquad')
                         AND s.session_date >= %s AND s.session_date <= %s""",
                    (owner, min(dates) - timedelta(days=45), max(dates)))
        tm = [dict(r) for r in cur.fetchall()]

    by_key, by_key_np = {}, {}
    for t in tm:
        cnt = f"{t['balls']}-{t['strikes']}"
        by_key.setdefault((_last(t["catcher"]), _last(t["pitcher"]), t["inning"], cnt), []).append(t)
        by_key_np.setdefault((_last(t["catcher"]), t["inning"], cnt), []).append(t)

    events, agg, by_pitcher = [], {}, {}
    for r in rows:
        r["session_date"] = r["session_date"].isoformat()
        c = agg.setdefault(r["catcher"], _new_agg(r["catcher"]))
        if r["game_no"]:
            c["games"].add((r["session_date"], r["game_no"]))
        o, e = r["outcome"] or "", r["event"] or ""
        if e == "Block" or o in BLOCK_OUTCOMES:
            c["block_opps"] += 1
            kept = o == "Block Success" or (e == "Block" and o in THROW_OUTCOMES)   # kept in front, runner still went
            if kept:
                c["block_success"] += 1
            elif o == "Block Miss":
                c["block_miss"] += 1
            elif o == "Passed Ball":
                c["pb"] += 1
            if r.get("two_strike"):
                c["k2_opps"] += 1
                c["k2_success"] += 1 if kept else 0
            if r.get("pitcher"):
                bp = by_pitcher.setdefault(r["pitcher"], {"pitcher": r["pitcher"], "opps": 0, "success": 0, "miss": 0, "pb": 0})
                bp["opps"] += 1
                bp["success"] += 1 if kept else 0
                bp["miss"] += 1 if o == "Block Miss" else 0
                bp["pb"] += 1 if o == "Passed Ball" else 0
        if o == "Stolen Base":
            c["sb"] += 1
        elif o == "Caught Stealing":
            c["cs"] += 1
        elif o == "Pickoff":
            c["pickoffs"] += 1
        if e == "Drop 3rd":
            c["drop3"] += 1
            c["drop3_outs"] += 1 if o == "Out" else 0
        if e == "Throwdown":
            c["throwdowns"] += 1
        if r["pop_time"] is not None:
            c["pops"].append(r["pop_time"])
        if r["grade"] is not None:
            c["grades"].append(r["grade"])
        if r["note"] and r["pop_time"] is None and e != "Throwdown":
            c["notes"].append(r["note"])
        match = None
        if e in ("Block", "Drop 3rd") and r["inning"] is not None and r["count"]:
            d0 = date.fromisoformat(r["session_date"])
            window = lambda t: 0 <= (d0 - t["session_date"]).days <= 45  # noqa: E731
            cands = [t for t in by_key.get((_last(r["catcher"]), _last(r.get("pitcher")), r["inning"], r["count"]), []) if window(t)] if r.get("pitcher") else []
            if not cands:
                cands = [t for t in by_key_np.get((_last(r["catcher"]), r["inning"], r["count"]), []) if window(t)]
            pref = [t for t in cands if t["pitch_call"] in ("BallinDirt", "BallCalled", "StrikeSwinging")]
            pool = pref or cands
            if len(pool) == 1:
                t = pool[0]
                match = {"pitch_id": t["id"], "date": t["session_date"].isoformat(), "ptype": t["ptype"],
                         "velo": round(float(t["rel_speed"]), 1) if t["rel_speed"] else None,
                         "pitcher": t["pitcher"], "call": t["pitch_call"],
                         "loc_side": round(float(t["plate_loc_side"]), 2) if t["plate_loc_side"] is not None else None,
                         "loc_height": round(float(t["plate_loc_height"]), 2) if t["plate_loc_height"] is not None else None}
        events.append({**r, "tm": match})

    catchers = []
    for c in agg.values():
        opps, steals = c["block_opps"], c["sb"] + c["cs"]
        catchers.append({
            "catcher": c["catcher"], "games": len(c["games"]),
            "block_opps": opps, "block_success": c["block_success"], "block_miss": c["block_miss"], "pb": c["pb"],
            "block_pct": round(100 * c["block_success"] / opps, 1) if opps else None,
            "k2_opps": c["k2_opps"], "k2_block_pct": round(100 * c["k2_success"] / c["k2_opps"], 1) if c["k2_opps"] else None,
            "sb": c["sb"], "cs": c["cs"], "cs_pct": round(100 * c["cs"] / steals, 1) if steals else None,
            "pickoffs": c["pickoffs"], "drop3": c["drop3"], "drop3_outs": c["drop3_outs"],
            "throwdowns": c["throwdowns"],
            "avg_pop": round(sum(c["pops"]) / len(c["pops"]), 2) if c["pops"] else None,
            "best_pop": round(min(c["pops"]), 2) if c["pops"] else None, "pops_n": len(c["pops"]),
            "avg_grade": round(sum(c["grades"]) / len(c["grades"]), 2) if c["grades"] else None, "grades_n": len(c["grades"]),
            "notes": c["notes"][:8],
        })
    catchers.sort(key=lambda c: -(c["block_opps"] + c["sb"] + c["cs"] + c["throwdowns"]))
    pitchers = sorted(by_pitcher.values(), key=lambda p: -p["opps"])
    for p in pitchers:
        p["block_pct"] = round(100 * p["success"] / p["opps"], 1) if p["opps"] else None
    return {"catchers": catchers, "by_pitcher": pitchers, "events": events, "dates": [d.isoformat() for d in dates]}


@router.delete("/trackman/catcher-log/{session_date}")
def catcher_log_delete(session_date: str, owner: str = Depends(_write_gate)):
    with get_connection() as conn:
        cur = conn.cursor()
        _ensure_table(cur)
        cur.execute("DELETE FROM tm_catcher_log WHERE owner_user_id = %s AND session_date = %s", (owner, session_date))
        n = cur.rowcount
        conn.commit()
    return {"status": "ok", "deleted": n}
